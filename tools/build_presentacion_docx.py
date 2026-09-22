#!/usr/bin/env python3
"""Convierte el guion de presentación a Word, con tipografía para leerlo de pie.

    PYTHONPATH=. python tools/build_presentacion_docx.py

El .docx no se versiona: se regenera desde el Markdown, que es la fuente. Así
no hay dos documentos que puedan decir cosas distintas —y las cifras del guion
están comprobadas contra el motor por tests/test_pension_feedback.py—.

Dos decisiones de formato que no son estéticas:

  · El cuerpo va a 12,5 puntos con interlineado 1,3. Esto se lee de pie, a
    distancia y mientras se habla, no sentado delante de la pantalla.

  · Las acotaciones («Pantalla: …») van en cursiva y en gris, sangradas. Son
    indicaciones escénicas, no texto que se diga en voz alta, y confundirlas
    delante del tribunal es justo lo que no puede pasar.
"""
from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
ORIGEN = RAIZ / "docs/PRESENTACION_10MIN.md"
DESTINO = RAIZ / "docs/PRESENTACION_10MIN.docx"


def _estilo(style, *, size=None, color=None, bold=None, italic=None,
            antes=None, despues=None, interlineado=None, fuente=None,
            sangria=None):
    f = style.font
    if size is not None:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if fuente is not None:
        f.name = fuente
    pf = style.paragraph_format
    if antes is not None:
        pf.space_before = Pt(antes)
    if despues is not None:
        pf.space_after = Pt(despues)
    if sangria is not None:
        pf.left_indent = Inches(sangria)
    if interlineado is not None:
        pf.line_spacing = interlineado
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def build(origen: Path = ORIGEN, destino: Path = DESTINO) -> Path:
    subprocess.run(
        ["pandoc", str(origen), "-o", str(destino), "--from", "gfm",
         "--toc", "--toc-depth=2", "--metadata", "toc-title=Índice"],
        check=True)

    d = Document(str(destino))
    S = d.styles
    _estilo(S["Normal"], size=12.5, interlineado=1.3, despues=9, fuente="Calibri")
    _estilo(S["Title"], size=24, bold=True, color="0F2B46", despues=14)
    _estilo(S["Heading 1"], size=19, bold=True, color="0F2B46", antes=22, despues=8)
    _estilo(S["Heading 2"], size=15.5, bold=True, color="087F8C", antes=20, despues=7)
    nombres = {s.name for s in S}
    for nombre in ("Block Text", "Quote", "Block Quote"):
        if nombre in nombres:
            _estilo(S[nombre], size=11.5, italic=True, color="5A6B7A",
                    antes=6, despues=10, sangria=0.28)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Inches(0.95)
        sec.top_margin = sec.bottom_margin = Inches(0.85)
    d.save(str(destino))
    return destino


if __name__ == "__main__":
    out = build()
    doc = Document(str(out))
    print(f"{out}  {out.stat().st_size / 1000:.1f} kB")
    print(f"  {len(doc.paragraphs)} párrafos · {len(doc.tables)} tablas · "
          f"{sum(1 for p in doc.paragraphs if p.style.name.startswith('Heading'))} encabezados")
